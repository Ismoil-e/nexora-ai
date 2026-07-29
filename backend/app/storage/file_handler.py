import io
import os
from typing import BinaryIO

from app.core.config import settings
from app.utils.exceptions import ValidationException


class FileHandler:
    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".csv", ".txt", ".png", ".jpg", ".jpeg", ".webp"}

    def validate_file(self, filename: str, file_size: int):
        ext = os.path.splitext(filename)[1].lower()
        if ext not in self.ALLOWED_EXTENSIONS:
            raise ValidationException(f"Ruxsat etilmagan fayl formati: {ext}")

        if file_size > settings.max_file_size_bytes:
            raise ValidationException(f"Fayl hajmi {settings.max_file_size_mb}MB dan katta")

    async def extract_text(self, file_data: bytes, filename: str) -> str:
        import pdfplumber
        from docx import Document
        import openpyxl
        import csv

        ext = os.path.splitext(filename)[1].lower()

        if ext == ".pdf":
            return await self._extract_pdf_text(file_data)
        elif ext in (".docx", ".doc"):
            return await self._extract_docx_text(file_data)
        elif ext in (".xlsx", ".csv"):
            return await self._extract_spreadsheet_text(file_data, ext)
        elif ext == ".txt":
            return file_data.decode("utf-8", errors="replace")
        else:
            return ""

    async def _extract_pdf_text(self, file_data: bytes) -> str:
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_data)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return "\n\n".join(text_parts)

    async def _extract_docx_text(self, file_data: bytes) -> str:
        doc = Document(io.BytesIO(file_data))
        return "\n".join([p.text for p in doc.paragraphs])

    async def _extract_spreadsheet_text(self, file_data: bytes, ext: str) -> str:
        text_parts = []

        if ext == ".csv":
            content = file_data.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(content))
            for row in reader:
                text_parts.append(" | ".join(row))
        else:
            wb = openpyxl.load_workbook(io.BytesIO(file_data), read_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                text_parts.append(f"Sheet: {sheet}")
                for row in ws.iter_rows(values_only=True):
                    text_parts.append(" | ".join([str(c) if c else "" for c in row]))
            wb.close()

        return "\n".join(text_parts)

    async def chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 200) -> list[str]:
        if not text:
            return []

        chunks = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = min(start + chunk_size, text_len)
            if end < text_len:
                last_space = text.rfind(" ", start, end)
                if last_space > start:
                    end = last_space

            chunks.append(text[start:end].strip())
            start = end - overlap if end < text_len else text_len

        return chunks
